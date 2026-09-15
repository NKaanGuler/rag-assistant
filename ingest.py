"""Build the RAG database: read docs, split into chunks, embed, store in sqlite.

Run this file directly (python ingest.py) whenever the files in the docs
folder change. It rebuilds the database from scratch each time, so the
database always matches exactly what is in the folder.
"""

import json
import os
import re
import sqlite3
from pathlib import Path

import rag

DOCS_DIR = "docs"
SUPPORTED_TYPES = (".txt", ".md", ".pdf", ".docx", ".pptx")


def read_document(path):
    """Return the plain text of a .txt / .md / .pdf / .docx / .pptx file.

    Pages, paragraphs and slides are joined with blank lines so that
    chunk_text() can treat them as paragraph boundaries.
    """
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        import logging
        from pypdf import PdfReader
        # Slide PDFs often trigger harmless "wrong pointing object" warnings;
        # keep them out of the console so the app's output stays readable.
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        pages = [page.extract_text() or "" for page in PdfReader(str(path)).pages]
        return "\n\n".join(pages)
    if suffix == ".docx":
        from docx import Document
        paragraphs = [p.text for p in Document(str(path)).paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation
        slides = []
        for slide in Presentation(str(path)).slides:
            texts = []
            for shape in slide.shapes:
                # Only shapes with a text frame (titles, bullet boxes) have text.
                text = getattr(shape, "text", "")
                if text and text.strip():
                    texts.append(text)
            if texts:
                slides.append("\n".join(texts))
        return "\n\n".join(slides)
    raise ValueError(f"Unsupported file type: {path.name}")


def split_long_paragraph(paragraph, max_chars):
    """Split one over-long paragraph at sentence ends into pieces <= max_chars.

    PDF and slide text often has no blank lines at all, so a whole page can
    arrive as one "paragraph". Cutting at sentence ends keeps each piece
    readable; a fixed-width cut is only the last resort for text with no
    sentence boundaries.
    """
    pieces = []
    current = ""
    for sentence in re.split(r"(?<=[.!?])\s+|\n", paragraph):
        sentence = sentence.strip()
        if not sentence:
            continue
        if len(sentence) > max_chars:
            if current:
                pieces.append(current)
                current = ""
            for start in range(0, len(sentence), max_chars):
                pieces.append(sentence[start:start + max_chars])
            continue
        if current and len(current) + 1 + len(sentence) > max_chars:
            pieces.append(current)
            current = sentence
        else:
            current = (current + " " + sentence) if current else sentence
    if current:
        pieces.append(current)
    return pieces


def chunk_text(text, max_chars=1200):
    """Split text into chunks of about max_chars, keeping paragraphs whole.

    We split on blank lines (paragraph breaks) and then greedily pack
    consecutive paragraphs into one chunk while it stays under max_chars.
    We never cut inside a normal paragraph, because half a paragraph is
    hard for the model (and for humans) to make sense of; only a paragraph
    that is itself longer than max_chars gets split, at sentence ends.
    """
    # A "paragraph" is a block of text separated by one or more blank lines.
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

    # Over-long paragraphs (typical for PDF pages) are split first.
    expanded = []
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            expanded.extend(split_long_paragraph(paragraph, max_chars))
        else:
            expanded.append(paragraph)
    paragraphs = expanded

    chunks = []
    current = ""
    for paragraph in paragraphs:
        if not current:
            # First paragraph of a new chunk — always take it, even if it is
            # longer than max_chars on its own.
            current = paragraph
        elif len(current) + 2 + len(paragraph) <= max_chars:
            # +2 accounts for the blank line we put back between paragraphs.
            current = current + "\n\n" + paragraph
        else:
            # Adding this paragraph would make the chunk too big:
            # close the current chunk and start a fresh one.
            chunks.append(current)
            current = paragraph
    if current:
        chunks.append(current)

    return chunks


def ingest(db_path=rag.DB_PATH, docs_dir=DOCS_DIR):
    """Read every supported document in docs_dir and store embedded chunks."""
    # Check the docs folder before doing any expensive model work, so a
    # missing folder gives a clear error instead of an empty database.
    docs_path = Path(docs_dir)
    if not docs_path.is_dir():
        raise SystemExit(
            f"Docs folder '{docs_dir}' not found. Create it and put your "
            f"documents ({', '.join(SUPPORTED_TYPES)}) inside, then run ingest again."
        )
    files = sorted(
        (p for p in docs_path.iterdir() if p.suffix.lower() in SUPPORTED_TYPES),
        key=lambda p: p.name.lower(),
    )
    if not files:
        raise SystemExit(
            f"No supported documents ({', '.join(SUPPORTED_TYPES)}) found in "
            f"'{docs_dir}'. Add some first, then run ingest again."
        )

    rag.get_manager()
    model = rag.load_model(rag.EMBEDDING_MODEL)
    embedding_client = model.get_embedding_client()

    # Build into a temporary file and swap it in only when everything
    # succeeded, so a failure half-way never destroys the existing database.
    tmp_path = db_path + ".building"
    Path(tmp_path).unlink(missing_ok=True)
    connection = sqlite3.connect(tmp_path)
    try:
        cursor = connection.cursor()

        # Drop and recreate the table so re-running ingest gives a clean
        # database instead of piling duplicate rows on top of the old ones.
        cursor.execute("DROP TABLE IF EXISTS chunks")
        cursor.execute(
            """
            CREATE TABLE chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source TEXT,
                chunk_index INTEGER,
                content TEXT,
                embedding TEXT
            )
            """
        )

        total = 0
        for path in files:
            text = read_document(path)
            chunks = chunk_text(text)
            if not chunks:
                print(f"{path.name}: empty file, skipped")
                continue

            # One batch call per document is much faster than embedding the
            # chunks one at a time.
            response = embedding_client.generate_embeddings(chunks)
            vectors = [item.embedding for item in response.data]

            for index, (content, vector) in enumerate(zip(chunks, vectors)):
                cursor.execute(
                    "INSERT INTO chunks (source, chunk_index, content, embedding) "
                    "VALUES (?, ?, ?, ?)",
                    # json.dumps turns the vector into text sqlite can store.
                    (path.name, index, content, json.dumps(vector)),
                )

            total += len(chunks)
            print(f"{path.name}: {len(chunks)} chunks")

        connection.commit()
    except BaseException:
        # Something failed mid-build: throw away the temporary file and keep
        # whatever database was there before.
        connection.close()
        Path(tmp_path).unlink(missing_ok=True)
        raise
    connection.close()
    os.replace(tmp_path, db_path)

    if total == 0:
        print("WARNING: all documents were empty - the database has no chunks.")
    print(f"Done. Stored {total} chunks in {db_path}.")

    # Free the memory the embedding model was using — whoever needs it next
    # (for example main.py) will load it again.
    model.unload()


if __name__ == "__main__":
    ingest()
